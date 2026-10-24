"""
Single Page Executive Overview Generator for BioSecure AI.
Generates:
1. BioSecure_AI_One_Page_Overview.pdf (Exactly 1 Page, professionally formatted & balanced)
2. BioSecure_AI_One_Page_Overview.docx (Matching Word document)
Strict constraints:
- Zero mentions of patents, IDF, or patent filing terms.
- Exact single-page layout with publication-grade spacing.
"""

import os
import fitz
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# Palette
NAVY = colors.HexColor('#0F172A')       # Slate 900
ROYAL = colors.HexColor('#1E3A8A')      # Blue 900
ACCENT = colors.HexColor('#2563EB')     # Blue 600
SKY = colors.HexColor('#0284C7')        # Sky 600
DARK = colors.HexColor('#1E293B')       # Slate 800
MUTED = colors.HexColor('#475569')      # Slate 600
LIGHT_BG = colors.HexColor('#F8FAFC')   # Slate 50
ALT_ROW = colors.HexColor('#F1F5F9')    # Slate 100
BORDER = colors.HexColor('#CBD5E1')     # Slate 300
LINE_GRAY = colors.HexColor('#E2E8F0')  # Slate 200
EMERALD = colors.HexColor('#059669')    # Green 600
AMBER = colors.HexColor('#D97706')      # Amber 600

PAGE_WIDTH, PAGE_HEIGHT = letter
MARGIN = 36  # 0.5 in
USABLE_WIDTH = PAGE_WIDTH - (2 * MARGIN)  # 540 pt


class SinglePageCanvas(canvas.Canvas):
    """Custom canvas for single page running header and footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)

    def showPage(self):
        # Header Rule
        self.saveState()
        self.setFont("Helvetica-Bold", 7.5)
        self.setFillColor(ROYAL)
        self.drawString(MARGIN, PAGE_HEIGHT - 23, "BIOSECURE AI")
        self.setFont("Helvetica", 7.5)
        self.setFillColor(MUTED)
        self.drawString(MARGIN + 62, PAGE_HEIGHT - 23, "|   EXECUTIVE SYSTEM OVERVIEW & ARCHITECTURAL BRIEF")
        self.drawRightString(PAGE_WIDTH - MARGIN, PAGE_HEIGHT - 23, "PRODUCTION RELEASE v3.2")
        self.setStrokeColor(LINE_GRAY)
        self.setLineWidth(0.7)
        self.line(MARGIN, PAGE_HEIGHT - 27, PAGE_WIDTH - MARGIN, PAGE_HEIGHT - 27)

        # Footer Rule
        self.setFont("Helvetica", 7.2)
        self.setFillColor(MUTED)
        self.drawString(MARGIN, 18, "BioSecure AI Web System  |  Institutional Biometric Systems Architecture  |  Production Release v3.2")
        self.setFont("Helvetica-Bold", 7.2)
        self.setFillColor(NAVY)
        self.drawRightString(PAGE_WIDTH - MARGIN, 18, "Executive One-Pager  |  Page 1 of 1")
        self.setStrokeColor(LINE_GRAY)
        self.setLineWidth(0.7)
        self.line(MARGIN, 26, PAGE_WIDTH - MARGIN, 26)
        self.restoreState()
        super().showPage()


def build_one_page_pdf(filename="BioSecure_AI_One_Page_Overview.pdf"):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=MARGIN,
        rightMargin=MARGIN,
        topMargin=32,
        bottomMargin=32
    )

    style_kicker = ParagraphStyle('K', fontName='Helvetica-Bold', fontSize=8.0, leading=10.0, textColor=SKY, spaceAfter=2)
    style_title = ParagraphStyle('T', fontName='Helvetica-Bold', fontSize=20.0, leading=23.0, textColor=NAVY, spaceAfter=2)
    style_subtitle = ParagraphStyle('S', fontName='Helvetica', fontSize=9.5, leading=12.5, textColor=ROYAL, spaceAfter=6)
    
    style_h2 = ParagraphStyle('H2', fontName='Helvetica-Bold', fontSize=9.5, leading=12.5, textColor=NAVY, spaceBefore=4.5, spaceAfter=2.5)
    style_body = ParagraphStyle('B', fontName='Helvetica', fontSize=8.2, leading=11.4, textColor=DARK, spaceAfter=3.5)
    
    style_callout_t = ParagraphStyle('CT', fontName='Helvetica-Bold', fontSize=8.8, leading=11.5, textColor=ACCENT, spaceAfter=2.5)
    style_callout_b = ParagraphStyle('CB', fontName='Helvetica', fontSize=8.0, leading=11.0, textColor=DARK)
    
    style_th = ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=7.6, leading=9.5, textColor=colors.white, alignment=0)
    style_td = ParagraphStyle('TD', fontName='Helvetica', fontSize=7.5, leading=9.8, textColor=DARK)
    style_td_b = ParagraphStyle('TDB', fontName='Helvetica-Bold', fontSize=7.5, leading=9.8, textColor=NAVY)

    story = []

    # Title Banner
    story.append(Paragraph("INSTITUTIONAL TECHNICAL SPECIFICATION & SYSTEM BRIEFING", style_kicker))
    story.append(Paragraph("BIOSECURE AI — SYSTEM OVERVIEW", style_title))
    story.append(Paragraph("Automated Multi-Face Recognition Attendance & Autonomous Biometric Drift Monitoring Platform", style_subtitle))

    # Core Purpose & Identity Box
    callout_content = [
        Paragraph("<b>PROJECT ESSENCE — WHAT THE PROJECT IS:</b>", style_callout_t),
        Paragraph(
            "<b>BioSecure AI</b> is an enterprise-grade, contactless attendance automation platform designed for universities, academic institutions, and organizations. It replaces manual roll calls (which waste 10–15 minutes of lecture time), eliminates magnetic/RFID swipe cards prone to proxy attendance ('buddy-punching'), and removes the physical queues and sanitary concerns of optical fingerprint scanners. By capturing or uploading a single classroom group photograph, the platform autonomously detects and aligns dozens of faces in parallel, extracts 512D deep embeddings, matches identities against enrolled student rosters in <b>under 1 millisecond on CPU</b>, logs attendance records to a PostgreSQL database, and proactively tracks biometric template aging across semesters.",
            style_callout_b
        )
    ]
    t_callout = Table([[callout_content]], colWidths=[USABLE_WIDTH])
    t_callout.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), LIGHT_BG),
        ('BOX', (0,0), (-1,-1), 0.5, BORDER),
        ('LINELEFT', (0,0), (0,0), 3.5, ACCENT),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_callout)
    story.append(Spacer(1, 5))

    # 5-Stage Pipeline Flow
    story.append(Paragraph("HOW IT WORKS — THE 5-STAGE AUTOMATED INFERENCE PIPELINE", style_h2))
    pipeline_data = [
        [
            Paragraph("<b>1. Multi-Face Capture</b><br/>Captures classroom photo via smartphone PWA window, webcam stream, or file upload.", style_td),
            Paragraph("<b>2. Detection & Align</b><br/>RetinaFace locates 50+ faces simultaneously, extracting 5 landmarks for canonical alignment.", style_td),
            Paragraph("<b>3. ArcFace Encoding</b><br/>Extracts 512-dimensional deep representations normalized to unit Euclidean sphere (||v|| = 1.0).", style_td),
            Paragraph("<b>4. In-Memory Match</b><br/>NumPy C-BLAS computes Q &middot; M^T batch dot product, matching full rosters in &lt;1.0 ms on CPU.", style_td),
            Paragraph("<b>5. Log & Drift Engine</b><br/>Logs 'Present' in PostgreSQL pgvector. Evaluates 3D Pose Gate and accumulates EWMA drift.", style_td),
        ]
    ]
    t_pipe = Table(pipeline_data, colWidths=[108, 108, 108, 108, 108])
    t_pipe.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), ALT_ROW),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 4.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4.5),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_pipe)
    story.append(Spacer(1, 5))

    # Core Capabilities Matrix (8 Comprehensive Rows)
    story.append(Paragraph("CORE CAPABILITIES & SYSTEM ARCHITECTURE", style_h2))
    caps_data = [
        [Paragraph("System Capability", style_th), Paragraph("Technical Implementation", style_th), Paragraph("Operational Impact & Performance Profile", style_th)],
        [
            Paragraph("<b>Dense Group Recognition</b>", style_td_b),
            Paragraph("RetinaFace ResNet50 + ArcFace Buffalo_L ONNX", style_td),
            Paragraph("Identifies 50+ students in parallel from a single wide-angle snapshot in &lt;900 ms turnaround.", style_td)
        ],
        [
            Paragraph("<b>Sub-Millisecond Matching</b>", style_td_b),
            Paragraph("NumPy BLAS Matrix Dot-Product (AVX2 SIMD)", style_td),
            Paragraph("&lt;1.0 ms CPU execution for 1,000 students; fallback to Supabase pgvector HNSW RPC query.", style_td)
        ],
        [
            Paragraph("<b>Biometric Drift Engine</b>", style_td_b),
            Paragraph("3D Pose Gate (|Yaw|&le;25&deg;, |Pitch|&le;20&deg;) + EWMA (&alpha;=0.30)", style_td),
            Paragraph("Continuously monitors facial template aging (hair, beard, weight) over multi-year degree programs.", style_td)
        ],
        [
            Paragraph("<b>Neural OCR ID Reader</b>", style_td_b),
            Paragraph("RapidOCR (DBNet+SVTR ONNX) + OpenCV Multi-Pass", style_td),
            Paragraph("Extracts student Name, Roll ID, Program, Branch, and Email directly from physical ID cards in &lt;1.2s.", style_td)
        ],
        [
            Paragraph("<b>Progressive Web App (PWA)</b>", style_td_b),
            Paragraph("Service Worker (sw.js) + HTML5 capture='environment'", style_td),
            Paragraph("Direct smartphone camera launch (ChatGPT-style), offline app shell fallback, and zero app-store friction.", style_td)
        ],
        [
            Paragraph("<b>Class Roster Modes</b>", style_td_b),
            Paragraph("Separate Single Class vs Combined Multi-Branch Mode", style_td),
            Paragraph("Enforces strict curriculum boundaries or evaluates joint auditorium cohorts dynamically.", style_td)
        ],
        [
            Paragraph("<b>Automated Alerting</b>", style_td_b),
            Paragraph("Asynchronous SMTP Email Dispatcher + /admin/drift", style_td),
            Paragraph("Dispatches warning emails upon CRITICAL drift (&ge;0.25) with 1-click re-enrollment template reset.", style_td)
        ],
        [
            Paragraph("<b>Security & Privacy (RLS)</b>", style_td_b),
            Paragraph("PostgreSQL Row-Level Security + Non-Invertible Vectors", style_td),
            Paragraph("Biometric embeddings cannot reverse-engineer faces; strict role separation (Instructor vs Admin).", style_td)
        ],
    ]
    t_caps = Table(caps_data, colWidths=[110, 190, 240])
    t_caps.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_ROW]),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 3.0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3.0),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_caps)
    story.append(Spacer(1, 5))

    # Technology Stack & Operational Benchmarks
    story.append(Paragraph("TECHNOLOGY STACK & OPERATIONAL BENCHMARKS", style_h2))
    bottom_data = [
        [
            Paragraph("<b>AI & Computer Vision:</b><br/>InsightFace (Buffalo_L), RetinaFace, ArcFace 512D, RapidOCR (DBNet+SVTR ONNX), OpenCV.", style_td),
            Paragraph("<b>Backend & Persistence:</b><br/>Python 3.10+, Flask Factory, Gunicorn WSGI, Nginx Proxy, PostgreSQL 16 + pgvector (HNSW Index).", style_td),
            Paragraph("<b>Client & Frontend:</b><br/>Progressive Web App (PWA), Service Worker, Tailwind CSS Dark Theme, Lucide Icons, WebRTC.", style_td),
            Paragraph("<b>Key Benchmarks:</b><br/>&bull; Vector Match: &lt;1.0 ms CPU<br/>&bull; End-to-End: &lt;1.0 s<br/>&bull; Lecture Time Saved: 15 min", style_td)
        ]
    ]
    t_bottom = Table(bottom_data, colWidths=[135, 145, 135, 125])
    t_bottom.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), LIGHT_BG),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 4.0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4.0),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_bottom)
    story.append(Spacer(1, 4))

    # Executive Conclusion Banner
    story.append(Paragraph(
        "<b>Summary Assessment:</b> BioSecure AI converts classroom attendance from an error-prone administrative hurdle into an ambient, instantaneous, and verifiable digital workflow. Combining sub-millisecond local vector matching with autonomous biometric drift tracking, the platform provides educational institutions with a high-accuracy, friction-free attendance solution.",
        style_body
    ))

    # Build document
    doc.build(story, canvasmaker=SinglePageCanvas)
    print(f"Successfully generated {filename}")


def build_one_page_docx(filename="BioSecure_AI_One_Page_Overview.docx"):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.5)
    sec.right_margin = Inches(0.5)

    # Header / Footer
    sec.header.paragraphs[0].text = "BIOSECURE AI  |  EXECUTIVE SYSTEM OVERVIEW & ARCHITECTURAL BRIEF"
    sec.footer.paragraphs[0].text = "BioSecure AI Web System  |  Institutional Technical Brief  |  Production Release v3.2"

    p_k = doc.add_paragraph()
    p_k.paragraph_format.space_before = Pt(0)
    p_k.paragraph_format.space_after = Pt(1)
    rk = p_k.add_run("INSTITUTIONAL TECHNICAL SPECIFICATION & SYSTEM BRIEFING")
    rk.font.name = "Calibri"
    rk.font.size = Pt(8.5)
    rk.font.bold = True
    rk.font.color.rgb = RGBColor(2, 132, 199)

    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(0)
    p_t.paragraph_format.space_after = Pt(2)
    rt = p_t.add_run("BIOSECURE AI — SYSTEM OVERVIEW")
    rt.font.name = "Calibri"
    rt.font.size = Pt(20)
    rt.font.bold = True
    rt.font.color.rgb = RGBColor(15, 23, 42)

    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_before = Pt(0)
    p_s.paragraph_format.space_after = Pt(6)
    rs = p_s.add_run("Automated Multi-Face Recognition Attendance & Autonomous Biometric Health Platform")
    rs.font.name = "Calibri"
    rs.font.size = Pt(10)
    rs.font.color.rgb = RGBColor(30, 58, 138)

    # Callout What It Is
    t_c = doc.add_table(rows=1, cols=1)
    t_c.columns[0].width = Inches(7.5)
    cell = t_c.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>'))
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    ct1 = cp.add_run("PROJECT ESSENCE — WHAT THE PROJECT IS:\n")
    ct1.font.name = "Calibri"
    ct1.font.size = Pt(9)
    ct1.font.bold = True
    ct1.font.color.rgb = RGBColor(37, 99, 235)
    ct2 = cp.add_run("BioSecure AI is an enterprise-grade, contactless attendance automation platform designed for universities, academic institutions, and organizations. It replaces manual roll calls (which waste 10–15 minutes of lecture time), eliminates magnetic/RFID swipe cards prone to proxy attendance ('buddy-punching'), and removes the physical queues and sanitary concerns of optical fingerprint scanners. By capturing or uploading a single classroom group photograph, the platform autonomously detects and aligns dozens of faces in parallel, extracts 512D deep embeddings, matches identities against enrolled student rosters in under 1 millisecond on CPU, logs attendance records to a PostgreSQL database, and proactively tracks biometric template aging across semesters.")
    ct2.font.name = "Calibri"
    ct2.font.size = Pt(8.5)
    ct2.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 5-Stage Pipeline
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(4)
    p_h1.paragraph_format.space_after = Pt(2)
    rh1 = p_h1.add_run("HOW IT WORKS — THE 5-STAGE AUTOMATED INFERENCE PIPELINE")
    rh1.font.name = "Calibri"
    rh1.font.size = Pt(10)
    rh1.font.bold = True
    rh1.font.color.rgb = RGBColor(15, 23, 42)

    t_pipe = doc.add_table(rows=1, cols=5)
    pipe_cols = [
        ("1. Capture", "Classroom photo via PWA mobile window, webcam, or upload."),
        ("2. Align", "RetinaFace locates 50+ faces, aligning 5 landmarks canonical."),
        ("3. Encode", "ArcFace deep network extracts 512D normalized vector."),
        ("4. Match", "In-memory BLAS dot-product Q · M^T matches in <1.0 ms CPU."),
        ("5. Log & Drift", "Logs 'Present' in pgvector; 3D Pose Gate updates EWMA drift.")
    ]
    for i, (title, desc) in enumerate(pipe_cols):
        cell = t_pipe.cell(0, i)
        cell.width = Inches(1.5)
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
        tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        pt1 = p.add_run(f"{title}\n")
        pt1.font.name = "Calibri"
        pt1.font.size = Pt(8.0)
        pt1.font.bold = True
        pt1.font.color.rgb = RGBColor(15, 23, 42)
        pt2 = p.add_run(desc)
        pt2.font.name = "Calibri"
        pt2.font.size = Pt(7.5)
        pt2.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Core Capabilities
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(4)
    p_h2.paragraph_format.space_after = Pt(2)
    rh2 = p_h2.add_run("CORE CAPABILITIES & SYSTEM ARCHITECTURE")
    rh2.font.name = "Calibri"
    rh2.font.size = Pt(10)
    rh2.font.bold = True
    rh2.font.color.rgb = RGBColor(15, 23, 42)

    caps_data = [
        ["System Capability", "Technical Implementation", "Operational Impact & Performance Profile"],
        ["Dense Group Recognition", "RetinaFace ResNet50 + ArcFace Buffalo_L ONNX", "Identifies 50+ students in parallel from a single wide-angle snapshot in <900 ms turnaround."],
        ["Sub-Millisecond Matching", "NumPy BLAS Matrix Dot-Product (AVX2 SIMD)", "<1.0 ms CPU execution for 1,000 students; fallback to Supabase pgvector HNSW RPC query."],
        ["Biometric Drift Engine", "3D Pose Gate (|Yaw|<=25°, |Pitch|<=20°) + EWMA (α=0.30)", "Continuously monitors facial template aging (hair, beard, weight) over multi-year degree programs."],
        ["Neural OCR ID Reader", "RapidOCR (DBNet+SVTR ONNX) + OpenCV Multi-Pass", "Extracts student Name, Roll ID, Program, Branch, and Email directly from physical ID cards in <1.2s."],
        ["Progressive Web App (PWA)", "Service Worker (sw.js) + HTML5 capture='environment'", "Direct smartphone camera launch (ChatGPT-style), offline app shell fallback, and zero app-store friction."],
        ["Class Roster Modes", "Separate Single Class vs Combined Multi-Branch Mode", "Enforces strict curriculum boundaries or evaluates joint auditorium cohorts dynamically."],
        ["Automated Alerting", "Asynchronous SMTP Email Dispatcher + /admin/drift", "Dispatches warning emails upon CRITICAL drift (>=0.25) with 1-click re-enrollment template reset."],
        ["Security & Privacy (RLS)", "PostgreSQL Row-Level Security + Non-Invertible Vectors", "Biometric embeddings cannot reverse-engineer faces; strict role separation (Instructor vs Admin)."]
    ]
    t_caps = doc.add_table(rows=len(caps_data), cols=3)
    t_caps.columns[0].width = Inches(1.5)
    t_caps.columns[1].width = Inches(2.7)
    t_caps.columns[2].width = Inches(3.3)
    for r_idx, row in enumerate(caps_data):
        for c_idx, val in enumerate(row):
            cell = t_caps.cell(r_idx, c_idx)
            tcPr = cell._tc.get_or_add_tcPr()
            bg = "0F172A" if r_idx == 0 else ("F1F5F9" if r_idx % 2 == 1 else "FFFFFF")
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>'))
            tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(7.8)
            if r_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)
                if c_idx == 0:
                    run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Tech Matrix
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(4)
    p_h3.paragraph_format.space_after = Pt(2)
    rh3 = p_h3.add_run("TECHNOLOGY STACK & OPERATIONAL BENCHMARKS")
    rh3.font.name = "Calibri"
    rh3.font.size = Pt(10)
    rh3.font.bold = True
    rh3.font.color.rgb = RGBColor(15, 23, 42)

    t_tech = doc.add_table(rows=1, cols=4)
    tech_cols = [
        ("AI & Computer Vision", "InsightFace (Buffalo_L), RetinaFace, ArcFace 512D, RapidOCR (DBNet+SVTR ONNX), OpenCV."),
        ("Backend & Database", "Python 3.10+, Flask, Gunicorn WSGI, Nginx Proxy, PostgreSQL 16 + pgvector (HNSW Index)."),
        ("Client & Frontend", "Progressive Web App (PWA), Service Worker (sw.js), Tailwind CSS, Lucide Icons, WebRTC."),
        ("Operational Value", "• Matching: <1.0 ms CPU\n• End-to-End: <1.0 s\n• Time Saved: 15 min/class")
    ]
    for i, (title, desc) in enumerate(tech_cols):
        cell = t_tech.cell(0, i)
        cell.width = Inches(1.87)
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
        tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        pt1 = p.add_run(f"{title}\n")
        pt1.font.name = "Calibri"
        pt1.font.size = Pt(8.0)
        pt1.font.bold = True
        pt1.font.color.rgb = RGBColor(15, 23, 42)
        pt2 = p.add_run(desc)
        pt2.font.name = "Calibri"
        pt2.font.size = Pt(7.5)
        pt2.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    p_end = doc.add_paragraph()
    re = p_end.add_run("Summary Assessment: BioSecure AI converts classroom attendance from an error-prone administrative hurdle into an ambient, instantaneous, and verifiable digital workflow. Combining sub-millisecond local vector matching with autonomous biometric drift tracking, the platform provides educational institutions with a high-accuracy, friction-free attendance solution.")
    re.font.name = "Calibri"
    re.font.size = Pt(8.0)
    re.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(filename)
    print(f"Successfully generated {filename}")


def verify_one_page_files(pdf_path="BioSecure_AI_One_Page_Overview.pdf", docx_path="BioSecure_AI_One_Page_Overview.docx"):
    # 1. PDF Verification
    doc_pdf = fitz.open(pdf_path)
    count = doc_pdf.page_count
    print(f"PDF Page Count: {count}")
    if count != 1:
        raise ValueError(f"Expected exactly 1 page for PDF, got {count}!")
    
    pdf_text = "".join(p.get_text() for p in doc_pdf).lower()
    banned = ['patent', 'idf', 'invention disclosure', 'prior art']
    for b in banned:
        if b in pdf_text:
            raise ValueError(f"Banned word '{b}' found in PDF!")
    print("One-Page PDF Verification PASSED: Exactly 1 page, 0 banned words.")

    # 2. DOCX Verification
    doc_word = docx.Document(docx_path)
    word_text = " ".join(p.text for p in doc_word.paragraphs)
    for t in doc_word.tables:
        for r in t.rows:
            for c in r.cells:
                word_text += " " + c.text
    word_text = word_text.lower()
    for b in banned:
        if b in word_text:
            raise ValueError(f"Banned word '{b}' found in DOCX!")
    print("One-Page DOCX Verification PASSED: 0 banned words.")


if __name__ == '__main__':
    build_one_page_pdf()
    build_one_page_docx()
    verify_one_page_files()
